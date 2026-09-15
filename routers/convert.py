"""Conversion tools — out of PDF and into PDF."""

from __future__ import annotations

import io
import os

import fitz
from fastapi import APIRouter, File, Form, UploadFile

from utils import (IMAGE_EXTS, bytes_response, cleanup, fail, file_response,
                   open_pdf, pdf_response, read_upload, require_module,
                   temp_path, zip_response)

router = APIRouter(prefix="/api", tags=["Convert"])

PAPER = {"a4": "a4", "letter": "letter", "legal": "legal", "a3": "a3"}


# --------------------------------------------------------------------------- #
# out of PDF
# --------------------------------------------------------------------------- #
@router.post("/pdf-to-jpg")
async def pdf_to_jpg(file: UploadFile = File(...), dpi: int = Form(150),
                     image_format: str = Form("jpg"), password: str = Form("")):
    image_format = image_format.lower()
    if image_format not in ("jpg", "png"):
        fail("Image format must be jpg or png.")
    dpi = max(72, min(dpi, 400))

    doc = open_pdf(await read_upload(file), password or None)
    try:
        entries = []
        for index in range(doc.page_count):
            pix = doc[index].get_pixmap(dpi=dpi)
            if image_format == "jpg":
                if pix.alpha:
                    pix = fitz.Pixmap(pix, 0)  # drop alpha, JPEG has none
                blob = pix.tobytes("jpeg", jpg_quality=88)
            else:
                blob = pix.tobytes("png")
            entries.append((f"page_{index + 1}.{image_format}", blob))
        return zip_response(entries, f"pdf_images_{image_format}.zip")
    finally:
        doc.close()


@router.post("/pdf-to-text")
async def pdf_to_text(file: UploadFile = File(...), password: str = Form("")):
    doc = open_pdf(await read_upload(file), password or None)
    try:
        chunks = []
        for index in range(doc.page_count):
            chunks.append(f"--- page {index + 1} ---\n{doc[index].get_text('text')}")
        text = "\n".join(chunks).strip()
        if not text.replace("-", "").strip():
            fail("No selectable text found. Run OCR on this file first.")
        return bytes_response(text.encode("utf-8"), "text/plain; charset=utf-8", "extracted.txt")
    finally:
        doc.close()


@router.post("/pdf-to-markdown")
async def pdf_to_markdown(file: UploadFile = File(...), password: str = Form("")):
    """Keeps headings, lists and tables via pymupdf4llm."""
    pymupdf4llm = require_module("pymupdf4llm", "pip install pymupdf4llm")
    doc = open_pdf(await read_upload(file), password or None)
    try:
        markdown = pymupdf4llm.to_markdown(doc)
        if not markdown.strip():
            fail("No text-based structure found. Run OCR on this file first.")
        return bytes_response(markdown.encode("utf-8"), "text/markdown; charset=utf-8",
                              "converted.md")
    finally:
        doc.close()


@router.post("/pdf-to-word")
async def pdf_to_word(file: UploadFile = File(...), password: str = Form("")):
    require_module("pdf2docx", "pip install pdf2docx")
    from pdf2docx import Converter

    data = await read_upload(file)
    open_pdf(data, password or None).close()  # validate before spending time

    src = temp_path(".pdf")
    dst = temp_path(".docx")
    with open(src, "wb") as fh:
        fh.write(data)
    try:
        converter = Converter(src, password=password or None)
        converter.convert(dst)
        converter.close()
    except Exception as exc:
        cleanup(src, dst)
        fail(f"Word conversion failed: {exc}")
    return file_response(
        dst,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "converted.docx", src)


@router.post("/pdf-to-excel")
async def pdf_to_excel(file: UploadFile = File(...), password: str = Form("")):
    """Every detected table becomes a worksheet."""
    require_module("openpyxl", "pip install openpyxl")
    from openpyxl import Workbook

    doc = open_pdf(await read_upload(file), password or None)
    try:
        book = Workbook()
        book.remove(book.active)
        found = 0
        for index in range(doc.page_count):
            for number, table in enumerate(doc[index].find_tables().tables, start=1):
                rows = table.extract()
                if not rows:
                    continue
                found += 1
                sheet = book.create_sheet(f"p{index + 1}_t{number}"[:31])
                for row in rows:
                    sheet.append(["" if cell is None else str(cell) for cell in row])
        if not found:
            fail("No tables were detected in this PDF.")
        buffer = io.BytesIO()
        book.save(buffer)
        return bytes_response(
            buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "tables.xlsx")
    finally:
        doc.close()


@router.post("/pdf-to-powerpoint")
async def pdf_to_powerpoint(file: UploadFile = File(...), dpi: int = Form(150),
                            password: str = Form("")):
    """One slide per page, rendered at full bleed."""
    require_module("pptx", "pip install python-pptx")
    from pptx import Presentation
    from pptx.util import Emu

    doc = open_pdf(await read_upload(file), password or None)
    try:
        first = doc[0].rect
        deck = Presentation()
        deck.slide_width = Emu(int(first.width / 72 * 914400))
        deck.slide_height = Emu(int(first.height / 72 * 914400))
        blank = deck.slide_layouts[6]
        for index in range(doc.page_count):
            pix = doc[index].get_pixmap(dpi=max(72, min(dpi, 300)))
            slide = deck.slides.add_slide(blank)
            slide.shapes.add_picture(io.BytesIO(pix.tobytes("png")), 0, 0,
                                     width=deck.slide_width, height=deck.slide_height)
        buffer = io.BytesIO()
        deck.save(buffer)
        return bytes_response(
            buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "slides.pptx")
    finally:
        doc.close()


# --------------------------------------------------------------------------- #
# into PDF
# --------------------------------------------------------------------------- #
@router.post("/images-to-pdf")
async def images_to_pdf(files: list[UploadFile] = File(...), page_size: str = Form("fit"),
                        margin: int = Form(0)):
    """'fit' makes each page match its image. Otherwise the image is centred on paper."""
    if not files:
        fail("Add at least one image.")
    out = fitz.open()
    try:
        for upload in files:
            blob = await read_upload(upload, IMAGE_EXTS)
            try:
                image = fitz.open(stream=blob, filetype=os.path.splitext(upload.filename)[1][1:])
                pdf_bytes = image.convert_to_pdf()
                image.close()
            except Exception:
                fail(f"{upload.filename} could not be read as an image.")
            single = fitz.open("pdf", pdf_bytes)
            if page_size == "fit":
                out.insert_pdf(single)
            else:
                paper = fitz.paper_rect(PAPER.get(page_size, "a4"))
                page = out.new_page(width=paper.width, height=paper.height)
                area = paper + (margin, margin, -margin, -margin)
                page.show_pdf_page(_fit(single[0].rect, area), single, 0)
            single.close()
        if out.page_count == 0:
            fail("No images could be converted.")
        return pdf_response(out, "images.pdf")
    except Exception:
        out.close()
        raise


@router.post("/html-to-pdf")
async def html_to_pdf(url: str = Form(""), html: str = Form(""),
                      page_size: str = Form("a4")):
    """Renders a web page or a raw HTML snippet with PyMuPDF's Story engine."""
    source = html.strip()
    if url.strip():
        if not url.startswith(("http://", "https://")):
            fail("Enter a full URL starting with http:// or https://")
        httpx = require_module("httpx", "pip install httpx")
        try:
            async with httpx.AsyncClient(timeout=25, follow_redirects=True,
                                         headers={"User-Agent": "pdfRED/1.0"}) as client:
                response = await client.get(url)
            response.raise_for_status()
            source = response.text
        except Exception as exc:
            fail(f"Could not fetch that page: {exc}")
    if not source:
        fail("Provide a URL or paste some HTML.")

    paper = fitz.paper_rect(PAPER.get(page_size, "a4"))
    area = paper + (40, 40, -40, -40)
    path = temp_path(".pdf")
    try:
        story = fitz.Story(html=source)
        writer = fitz.DocumentWriter(path)
        more = 1
        guard = 0
        while more and guard < 500:
            device = writer.begin_page(paper)
            more, _ = story.place(area)
            story.draw(device)
            writer.end_page()
            guard += 1
        writer.close()
    except Exception as exc:
        cleanup(path)
        fail(f"That HTML could not be rendered: {exc}")
    return file_response(path, "application/pdf", "webpage.pdf")


# --------------------------------------------------------------------------- #
def _fit(source: fitz.Rect, target: fitz.Rect) -> fitz.Rect:
    scale = min(target.width / source.width, target.height / source.height)
    width, height = source.width * scale, source.height * scale
    x = target.x0 + (target.width - width) / 2
    y = target.y0 + (target.height - height) / 2
    return fitz.Rect(x, y, x + width, y + height)
